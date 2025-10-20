
import streamlit as st
import pandas as pd
import fitz  # PyMuPDF
from datetime import datetime, timedelta
from openpyxl import load_workbook
import os

# Função para extrair texto de PDF
def extrair_texto_pdf(caminho_pdf):
    texto = ""
    with fitz.open(caminho_pdf) as doc:
        for pagina in doc:
            texto += pagina.get_text()
    return texto

# Função para simular extração de dados do contrato e OIS (exemplo simplificado)
def extrair_dados(texto_contrato, texto_ois):
    return {
        "Número": "4500089999",
        "Objeto": "Exemplo de objeto extraído do contrato",
        "Contratada": "Empresa Exemplo Ltda",
        "CNPJ": "00.000.000/0001-00",
        "Assinatura": datetime(2025, 10, 1),
        "OIS": datetime(2025, 10, 5),
        "Marcos": [
            ("Entrega de documentação técnica", datetime(2025, 10, 10), ""),
            ("Indicação de responsável técnico", datetime(2025, 10, 7), "")
        ]
    }

# Função para atualizar planilha
def atualizar_planilha(dados, caminho_planilha):
    wb = load_workbook(caminho_planilha)
    if dados["Número"] in wb.sheetnames:
        ws = wb[dados["Número"]]
    else:
        ws = wb.create_sheet(title=dados["Número"])

    ws["A1"] = "Número do Contrato"
    ws["B1"] = dados["Número"]
    ws["A2"] = "Objeto do Contrato"
    ws["B2"] = dados["Objeto"]
    ws["A3"] = "Contratada"
    ws["B3"] = dados["Contratada"]
    ws["A4"] = "CNPJ"
    ws["B4"] = dados["CNPJ"]
    ws["A5"] = "Data de Assinatura do Contrato"
    ws["B5"] = dados["Assinatura"].strftime("%d/%m/%Y")
    ws["A6"] = "Data da OIS"
    ws["B6"] = dados["OIS"].strftime("%d/%m/%Y")

    ws.append(["Marco Contratual", "Prazo", "Data de Cumprimento"])
    for descricao, prazo, cumprimento in dados["Marcos"]:
        ws.append([descricao, prazo.strftime("%d/%m/%Y"), cumprimento])

    wb.save(caminho_planilha)

# Interface Streamlit
st.title("Agente IA para Gestão de Contratos")

caminho_contrato = st.text_input("Caminho do arquivo PDF do Contrato")
caminho_ois = st.text_input("Caminho do arquivo PDF da Ordem de Início de Serviços")
caminho_planilha = st.text_input("Caminho da planilha de controle (Excel)")

if st.button("Atualizar Planilha"):
    if os.path.exists(caminho_contrato) and os.path.exists(caminho_ois) and os.path.exists(caminho_planilha):
        texto_contrato = extrair_texto_pdf(caminho_contrato)
        texto_ois = extrair_texto_pdf(caminho_ois)
        dados = extrair_dados(texto_contrato, texto_ois)
        atualizar_planilha(dados, caminho_planilha)
        st.success(f"Contrato {dados['Número']} atualizado com sucesso na planilha.")
    else:
        st.error("Verifique os caminhos informados.")

# Alerta visual de prazos
if caminho_planilha and os.path.exists(caminho_planilha):
    wb = load_workbook(caminho_planilha)
    for nome in wb.sheetnames:
        ws = wb[nome]
        st.subheader(f"Contrato {nome}")
        for row in ws.iter_rows(min_row=7, values_only=True):
            if row[1] and isinstance(row[1], str):
                try:
                    prazo = datetime.strptime(row[1], "%d/%m/%Y")
                    dias_restantes = (prazo - datetime.today()).days
                    if dias_restantes < 0:
                        st.error(f"⚠️ {row[0]} - Prazo vencido em {row[1]}")
                    elif dias_restantes == 0:
                        st.warning(f"⚠️ {row[0]} - Prazo vence hoje ({row[1]})")
                    elif dias_restantes <= 2:
                        st.info(f"F514 {row[0]} - Prazo em {row[1]} ({dias_restantes} dias restantes)")
                except:
                    continue
